# -*- coding: utf-8 -*-

import asyncio
import re
import os
import websockets
import uuid
import json
from tencentcloud.common import credential
from tencentcloud.common.profile.client_profile import ClientProfile
from tencentcloud.common.profile.http_profile import HttpProfile
from tencentcloud.common.exception.tencent_cloud_sdk_exception import TencentCloudSDKException
from tencentcloud.lke.v20231130 import lke_client, models
import ssl
import certifi
import pymysql
import random
import tempfile
from pdfminer.high_level import extract_text
import subprocess
import shutil



ssl_context = ssl.create_default_context()
ssl_context.load_verify_locations(certifi.where())

bot_app_key = ""  # 机器人密钥，不是BotBizId (从运营接口人处获取)
visitor_biz_id = ""  # 访客 ID（外部系统提供，需确认不同的访客使用不同的 ID）
secret_id = ""  # 填入腾讯云AKSK密钥(从腾讯云控制台获取)
secret_key = ""  # 填入腾讯云AKSK密钥(从腾讯云控制台获取)
tencent_cloud_domain = ""  # 腾讯云API域名
scheme = "https"  # 协议
req_method = "POST"  # 请求方法
region = "ap-guangzhou"  # 地域
conn_type_api = 5  # API 访客


def get_session():
    # uuid
    return str(uuid.uuid1())


def get_request_id():
    return str(uuid.uuid1())


def api_token():
    try:
        # 实例化一个认证对象，入参需要传入腾讯云账户 SecretId 和 SecretKey，此处还需注意密钥对的保密
        # 代码泄露可能会导致 SecretId 和 SecretKey 泄露，并威胁账号下所有资源的安全性。以下代码示例仅供参考，建议采用更安全的方式来使用密钥，请参见：https://cloud.tencent.com/document/product/1278/85305
        # 密钥可前往官网控制台 https://console.cloud.tencent.com/cam/capi 进行获取
        cred = credential.Credential(secret_id, secret_key)
        # 实例化一个http选项，可选的，没有特殊需求可以跳过
        httpProfile = HttpProfile()
        httpProfile.endpoint = "lke.tencentcloudapi.com"

        # 实例化一个client选项，可选的，没有特殊需求可以跳过
        clientProfile = ClientProfile()
        clientProfile.httpProfile = httpProfile
        # 实例化要请求产品的client对象,clientProfile是可选的
        client = lke_client.LkeClient(cred, region, clientProfile)

        # 实例化一个请求对象,每个接口都会对应一个request对象
        req = models.GetWsTokenRequest()
        params = {
            "Type": conn_type_api,
            "BotAppKey": bot_app_key,
            "VisitorBizId": visitor_biz_id
        }
        req.from_json_string(json.dumps(params))
        # 返回的resp是一个GetWsTokenResponse的实例，与请求对象对应
        resp = client.GetWsToken(req)
        # 输出json格式的字符串回包
        print(resp.to_json_string())
        return resp.Token
    except TencentCloudSDKException as err:
        print(err)


pattern = r'\d+(.*)'
# ========================================================================
# ===============================模拟上传文档===============================
# ========================================================================
from docx import Document

def doc_receive(doc_path):

    #doc_path = "/Users/t200051/Desktop/慧心/新3份/摩优耐新材料（广西）有限公司碳刷生产加工项目环境影响报告表（报批稿）.docx"

    doc = Document(doc_path)

    # ===== 读取所有文字，包括段落和表格 =====
    full_text = []

    # 普通段落
    for para in doc.paragraphs:
        full_text.append(para.text.strip())

    # 表格里的内容
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text.strip())

    # 拼接成完整文本
    combined_text = "\n".join(full_text)
    #print("📄 提取的全文内容:\n", combined_text[:500])

    # ===== 正则提取建设单位和建设地点 =====
    company_match = re.search(r"建设单位.*?[：:]\s*(.*?)(?:[；。.\n\r]|$)", combined_text)
    location_match = re.search(r"建设地点[：:]\s*([^\n；。]+)", combined_text)
    projectname_match = re.search(r"项目名称[：:]\s*(.+?)[；。\n]", combined_text)

    company = company_match.group(1).strip() if company_match else None
    location = location_match.group(1).strip() if location_match else None
    projectname = projectname_match.group(1).strip() if projectname_match else None

    print("✅ 建设单位:", company)
    print("✅ 建设地点:", location)
    print("✅ 项目名称:", projectname)
    return location, company, projectname  #location, company = extract_info_from_doc(doc_path)
    #filled_msgs = fill_placeholders(messages, location, company)
#doc_receive(doc_path)

# file_id = random.randint(1000, 9999)
async def call_ws(token, message, conn, file_id, sort, title_text, doc_path):
    url = f"wss://wss.lke.cloud.tencent.com/v1/qbot/chat/conn/?EIO=4&transport=websocket"
    async with websockets.connect(url, ssl=ssl_context, ping_interval=20, ping_timeout=10) as ws:
        try:
            response = await asyncio.wait_for(ws.recv(), timeout=10)
        except asyncio.TimeoutError:
            print("⚠️ 等待服务端回应超时！")
            return
        response = await ws.recv()
        print("连接建立：" + response)

        auth = {"token": token}
        auth_message = f"40{json.dumps(auth)}"
        await ws.send(auth_message)

        # 再次接收（也建议加超时）
        try:
            response = await asyncio.wait_for(ws.recv(), timeout=10)
        except asyncio.TimeoutError:
            print("⚠️ 鉴权回应超时")
            return
        # 收到服务端鉴权后的包
        response = await ws.recv()
        print("鉴权结果：" + str(response))

        session_id = get_session()
        request_id = get_request_id()
        payload = {
            "payload": {
                "request_id": request_id,
                "session_id": session_id,
                "content": f"{message}",
            }
        }
        req_data = ["send", payload]

        send_data = f"42{json.dumps(req_data, ensure_ascii=False)}"
        print("发送数据：" + send_data)
        await ws.send(send_data)
        reply_text = ""  # 初始化 reply_text

        while True:
            rsp = await ws.recv()
            if rsp == '2':
                await ws.send("3")
                continue

            rsp_re_result = re.search(pattern, rsp).group(1)
            rsp_dict = json.loads(rsp_re_result)

            if rsp_dict[0] == "error":
                print(rsp_dict)
                break
            elif rsp_dict[0] == "reply":
                content = rsp_dict[1]["payload"]["content"]
                if rsp_dict[1]["payload"]["is_from_self"]:
                    print(f"from self: {content}")
                    continue
                else:

                    if rsp_dict[1]["payload"]["is_final"]:


                        payload = rsp_dict[1]["payload"]
                        final_content = payload["content"]
                        print(f"✅ 最终回复（{title_text}）：\n{final_content}\n")
                        # ✅ 插入数据库
                        title = f"自动生成段落 {sort}"
                        # file_id = random.randint(1000, 9999)
                        with conn.cursor() as cur:
                            location, company, projectname = doc_receive(doc_path)
                            sql = """
                                INSERT INTO ai_examine (file_id, title, reply, file_name, sort)
                                VALUES (%s, %s, %s, %s, %s)
                            """
                            doc_receive(doc_path)
                            cur.execute(sql, (file_id, title, final_content, projectname, 1))
                            conn.commit()

                        break

def run_websocket_processing(doc_path):
    file_id = random.randint(1000, 9999)
    conn = pymysql.connect(
        host='localhost',
        user='root',
        password='',
        database='ai_examine',
        charset='utf8mb4',
        cursorclass=pymysql.cursors.DictCursor,
        unix_socket='/tmp/mysql.sock'
    )
    # 🔥 清空 ai_examine 表
    with conn.cursor() as cur:
        cur.execute("DELETE FROM ai_examine;")
        conn.commit()
        print("✅ 数据表 ai_examine 已清空")
    location, company, projectname = doc_receive(doc_path)
    messages_template=["你是一个设计咨询有限公司的第三方评估员 ，只根据{{location}}的{{projectname}}总结一个：一、项目概况 ，请你只基于我提供的这篇文章内容进行总结，不得参考任何其他信息来源或自行补充内容。 如果原文中没有的信息，请不要编造。公文风格，段落格式输出，只写一段，以时间线介绍，不要总结字数。首先（一）根据建设项目基本情况介绍项目建设地点，然后根据建设项目工程分析的建设内容介绍建设内容。接着（二）主要保护目标，根据环境保护目标来写（本项目周边500m范围内大气环境，地表水环境，地下水环境，声环境，生态环境：即本项目周边500m范围内是否涉及饮用水源保护区、自然保护区、文物保护、风景名胜区等环境敏感区域。）永远用具体内容代替xx。你输出的批复状况，现有建设，环境影响评价全都不符合{{company}}的内容，一定要在{{company}}的文章里检索正确的信息",
              "你是一个设计咨询有限公司的第三方评估员 ，只根据{{location}}的{{projectname}}接着写二、执行标准（一）环境质量标准：根据五、环境保护措施监督检查清单下的表格的执行标准这一列来写。要素包括：大气环境、地表水环境、声环境、固体废物、土壤及地下水污染防治措施。（二）、污染物排放控制标准：根据三、区域环境质量现状、环境保护目标及评价标准的污染物排放控制标准来写（废气、废水、噪声、固体废物），不要总结字数。不要标注出处。公文格式（1000字）",
              "你是一个设计咨询有限公司的第三方评估员 ，只根据{{location}}的{{projectname}}接着写，三、区域环境质量现状，根据三、区域环境质量现状、环境保护目标及评价标准1.环境空气质量现状2地表水环境质量现状.3.声环境质量现状4.地下水、土壤5.生态环境质量现状来写。每个点用（中文数字）做标题，公文格式，1000字，不要总结字数。",
              "你是一个设计咨询有限公司的第三方评估员 ，只根据{{location}}的{{projectname}}接着写，四、主要环境影响及环境保护措施：根据文章的四、主要环境影响和保护措施来写，分施工期和运营期来写。每个期间都包括大气、水、声、固体废物、环境风险。公文段落格式，2000字，不需要结论和字数",
              "你是一个设计咨询有限公司的第三方评估员 ，只根据{{location}}的{{projectname}}接着写，五、评估结论：给出分析项目是否符合当地土地利用规划要求，产业发展政策及行业相关规范。从环保角度分析，项目的建设是否是可行的。只写结论，不需要提供依据，公文段落格式，不需要标注字数。"]
    messages = [
    msg.replace("{{location}}", location).replace("{{company}}", company).replace("{{projectname}}", projectname)
    for msg in messages_template
]

    print("\n📝 填充后的消息内容：\n", messages)
    for idx, msg in enumerate(messages, start=1):
        title_text = f"自动生成段落 {idx}"

        message = msg
        print(f"请求内容: {message}")
        token = api_token()
        print(token)
        if token == "":
            print("get token error")
            exit(0)
        asyncio.run(call_ws(token, msg, conn, file_id, idx, title_text, doc_path))
    return file_id


# if __name__ == "__main__":
#     run_websocket_processing()