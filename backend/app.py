from flask import Flask, send_file, jsonify, request, redirect, render_template, send_from_directory, g
from flask_cors import CORS
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from datetime import datetime
import io
import pymysql
import sys
import os
import platform
import subprocess
import threading


UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)  # 自动创建这个目录（如果不存在）

target_path = os.path.join(os.path.dirname(__file__), 'chat-demo-client', 'ws_client')
print("🔍 添加模块路径：", target_path)
sys.path.append(target_path)
from example_websocket import run_websocket_processing

# app = Flask(__name__, static_folder='static', static_url_path='/')
app = Flask(__name__)
CORS(app)
# 数据库连接信息
# db_config = {
#     'host': 'localhost',
#     'user': 'root',
#     'password': '',
#     'database': 'ai_examine',
#     'charset': 'utf8mb4',
#     'cursorclass': pymysql.cursors.DictCursor,
#     # 'unix_socket': '/tmp/mysql.sock'
# }
# # 仅在 Linux/macOS 上使用 unix_socket
# if platform.system() != 'Windows':
#     db_config['unix_socket'] = '/tmp/mysql.sock'

def get_connection():
    if 'db' not in g:
        config = {
            'host': 'localhost',
            'user': 'root',
            'password': '',
            'database': 'ai_examine',
            'charset': 'utf8mb4',
            'cursorclass': pymysql.cursors.DictCursor
        }
        if platform.system() != 'Windows':
            config['unix_socket'] = '/tmp/mysql.sock'
        g.db = pymysql.connect(**config)
    return g.db

@app.teardown_appcontext
def close_connection(exception):
    db = g.pop('db', None)
    if db is not None:
        db.close()
# conn = get_connection()
# 测试数据库连接
# try:
#     conn = get_connection()
#     #conn = pymysql.connect(**db_config)
#     print("✅ pymysql 连接成功")
# except Exception as e:
#     print(f"❌ pymysql 连接失败: {e}")
#     conn = None

# cur = conn.cursor()
# cur.execute("SELECT file_id, title, reply, file_name FROM ai_examine WHERE file_id = 1 ORDER BY sort")
# data = cur.fetchall()
#
# for row in data:
#     print(f"title: {row['title']}")
#     print(f"reply:\n{row['reply']}")
def get_news_data(file_id=None):
    try:
        conn = get_connection()
        cur = conn.cursor()
        if file_id:
            cur.execute("""
                SELECT file_id, title, reply, file_name
                FROM ai_examine
                WHERE file_id = %s
                ORDER BY sort
            """, (file_id,))
        else:
            cur.execute("""
                SELECT file_id, file_name
                FROM ai_examine
                GROUP BY file_id, file_name
                ORDER BY MIN(sort)
            """)
        data = cur.fetchall()
        cur.close()
        return data
    except Exception as e:
        raise Exception(f"数据库查询失败: {e}")



print(">>> 开始写入段落")
def generate_news_document(data):
    print("🎯 进入 generate_news_document()", flush=True)
    print(f"收到数据条数：{len(data)}", flush=True)

    print(f"[DEBUG] generate_news_document 收到 {len(data)} 条记录", file=sys.stderr)

    for i, row in enumerate(data):
        print(f"[DEBUG] 第 {i+1} 条记录: title = {row.get('title')}, reply 长度 = {len(row.get('reply', ''))}", flush=True)
        print(f"[DEBUG] reply 前 100 字：{row.get('reply', '')[:100]}", flush=True)
    """生成段落式文档（优化换行处理）"""
    doc = Document()

    # ===== 1. 设置全局样式 =====
    set_document_styles(doc)

    # ===== 2. 添加文档标题 =====
    # 获取第一条数据的file_name（假设所有数据的file_name相同）
    file_name = data[0]['file_name'] if data and 'file_name' in data[0] else None
    add_document_header(doc,file_name)

    # ===== 3. 添加内容（关键修改部分） =====
    add_news_paragraphs_optimized(doc, data)  # 改用优化后的函数

    # ===== 4. 添加文档页脚 =====
    add_document_footer(doc)

    # 保存到内存流
    file_stream = io.BytesIO()
    doc.save(file_stream)
    print("保存成功，文件大小（字节）", file_stream.tell())
    print("🧪 写入段落总数：", len(doc.paragraphs))
    for i, para in enumerate(doc.paragraphs, 1):
        print(f"第{i}段内容：{para.text[:50]}")
    file_stream.seek(0)

    return file_stream
    # from docx import Document
    # import io
    #
    # doc = Document()
    #
    # # 添加一个简单标题
    # doc.add_heading("这是一个测试标题", level=1)
    #
    # # 写入几段模拟内容
    # for i in range(1, 4):
    #     doc.add_paragraph(f"这是第{i}段内容，用于测试 Word 写入是否成功。")
    #
    # file_stream = io.BytesIO()
    # doc.save(file_stream)
    # file_stream.seek(0)
    # print("✅ 简化写入成功")
    # return file_stream
print(">>> 段落写入完成")
print(f"当前执行文件是：{__file__}")

def add_news_paragraphs_optimized(doc, data):
    print(f"共有 {len(data)} 条数据要写入 Word", flush=True)

    # 创建无间隔段落样式（如已存在则捕获异常）
    try:
        no_space_style = doc.styles.add_style('NoSpace', WD_STYLE_TYPE.PARAGRAPH)
        no_space_style.paragraph_format.space_after = Pt(0)
        no_space_style.paragraph_format.space_before = Pt(0)
        no_space_style.font.name = '微软雅黑'
        no_space_style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    except ValueError:
        pass  # 样式已存在

    for idx, ai_examine in enumerate(data, 1):
        title = ai_examine.get('title', f'未命名标题{idx}')
        reply = ai_examine['reply'].strip()

        doc.add_heading(title, level=2)
        doc.add_paragraph(reply)
        if not reply:
            print(f"[⚠️ 警告] 第 {idx} 条 reply 为空！", flush=True)
            doc.add_paragraph("(此段暂无内容)", style='NoSpace')
            continue

        print(f"[✅ 写入] 第 {idx} 条记录 - 标题: {title}", flush=True)
        print(f"[内容预览] {reply[:100]}", flush=True)

        # 写入 reply，每行为一段
        # for line in reply.splitlines():
        #     line = line.strip()
        #     if line:



        # ---- 分割线 ----
        # if idx < len(data):
        #     doc.add_paragraph("_" * 50, style='NoSpace').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def set_document_styles(doc):
    """设置文档基础样式（保持原有样式）"""
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

def add_document_header(doc,file_name=None):
    """文档标题页（保持原有样式）"""
    # title = doc.add_heading('AI辅助审查报告表', level=0)
    # 使用传入的file_name作为标题，如果没有则使用默认标题
    title_text = file_name if file_name else 'AI辅助审查报告表'

    title = doc.add_heading(title_text, level=0)
    # title = doc.add_heading('关于年产30万立方米加气混凝土切块技改项目环境影响报告表技术评估意见的函', level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x2A, 0x5C, 0xAA)
    title.runs[0].font.size = Pt(22)

    doc.add_paragraph("\n")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = subtitle.add_run(f"生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}")
    run.font.size = Pt(12)
    run.italic = True
    doc.add_page_break()

# def add_news_paragraphs(doc, data):
#     """按段落添加每条新闻"""
#     # doc.add_heading('内容', level=1)
#
#     for ai_examine in data:
#         # ---- 标题 ----
#         title_para = doc.add_paragraph()
#         title_para.paragraph_format.space_before = Pt(12)  # 段前间距
#         title_run = title_para.add_run(f"{ai_examine['title']}")
#         title_run.bold = True
#         title_run.font.size = Pt(14)
#         title_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)  # 深灰色
#
#
#         # ---- 正文 ----
#         content_para = doc.add_paragraph()
#         # content_para.paragraph_format.first_line_indent = Inches(0.5)  # 首行缩进
#         content_run = content_para.add_run(ai_examine['reply'])
#         content_run.font.size = Pt(12)
#
#         # 添加分割线（除最后一条外） 暂时无需分割线
#         # if news != data[-1]:
#         #     doc.add_paragraph("_" * 50).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

def add_document_footer(doc):
    """文档页脚（保持原有样式）"""
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.text = f""
    footer_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    footer_para.runs[0].font.size = Pt(9)
#导出所有数据 暂时不用
# @app.route('/export/news-paragraphs', methods=['GET'])
# @app.route('/export/auditfile', methods=['GET'])
# def export_news_paragraphs():
#     try:
#         data = get_news_data()
#         if not data:
#             return {"error": "没有找到数据"}, 404
        
#         word_file = generate_news_document(data)
#         return send_file(
#             word_file,
#             as_attachment=True,
#             download_name=f"报告_{datetime.now().strftime('%Y%m%d')}.docx",
#             mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
#         )
#     except Exception as e:
#         return {"error": str(e)}, 500
# 修改路由以接收file_id参数

# for para in data:  # 从 1 开始计数
#     print('wode', para['reply'].strip(''))
# print(data[3]['reply'])

# @app.route('/')
# @app.route('/upload')
# def serve_vue():
#     return send_from_directory('static', 'index.html')

# @app.route('/assets/<path:filename>')
# def serve_assets(filename):
#     return send_from_directory('static/assets', filename)

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


@app.route('/upload/doc', methods=['POST'])
def upload_doc():
    try:

        file = request.files.get('file')
        if not file:
            return jsonify({'error': '未提供文件'}), 400

        save_path = os.path.join(UPLOAD_FOLDER, 'uploaded.docx')
        filename = file.filename
        ext = os.path.splitext(filename)[1].lower()

        if ext not in ['.docx', '.doc', '.pdf']:
            return jsonify({'error': f'暂不支持的文件类型：{ext}，仅支持 .docx/.doc/.pdf'}), 400

        temp_path = os.path.join(UPLOAD_FOLDER, f"uploaded{ext}")
        file.save(temp_path)
        print("✅ 文件已保存为：", temp_path)

        # 最终统一生成为 uploaded.docx 以供后续处理
        final_path = os.path.join(UPLOAD_FOLDER, "uploaded.docx")

        # PDF → DOCX
        if ext == '.pdf':
            from pdfminer.high_level import extract_text
            from docx import Document

            text = extract_text(temp_path)

            # ✅ 清除 NULL 字节和不可打印字符
            safe_text = ''.join(c for c in text if c.isprintable() and ord(c) >= 32)

            doc = Document()

            # 按换行分段（保持段落结构）
            for line in safe_text.splitlines():
                line = line.strip()
                if line:  # 避免写入空行
                    doc.add_paragraph(line)

            doc.save(final_path)
            print("✅ PDF 已转换为 DOCX（保持多段落）")

        # DOC → DOCX

        elif ext == '.doc':
            # import subprocess

            # 1. 调用 libreoffice 转换
            subprocess.run([
                "libreoffice", "--headless", "--convert-to", "docx",
                "--outdir", UPLOAD_FOLDER, temp_path
            ], check=True)

            # 2. 获取实际生成的文件路径（uploaded.docx）
            base_name = os.path.splitext(os.path.basename(temp_path))[0]  # uploaded
            converted = os.path.join(UPLOAD_FOLDER, f"{base_name}.docx")
            print(f"✅ DOC 转换输出文件路径: {converted}")

            # 3. 替换最终统一路径
            if converted != final_path:
                os.replace(converted, final_path)
                print("✅ 替换为统一 DOCX 文件路径：", final_path)
        # DOCX 直接使用
        elif ext == '.docx':
            final_path = temp_path

        # 最终处理统一的 DOCX 文件
        # run_websocket_processing(final_path)
        threading.Thread(target=run_websocket_processing, args=(final_path,)).start()
        file_id = run_websocket_processing(final_path)
        return jsonify({'message': '上传成功', 'redirect_url': f"/export/auditfile/{file_id}"})
        print("✅ 上传处理完毕，准备跳转")
        #file.save(save_path)
        #print("✅ 已保存文件:", save_path)

        #run_websocket_processing(save_path)

        # return jsonify({'message': '上传成功', 'redirect_url': "/export/auditfile/1"})

    except Exception as e:
        print(f"❌ 后端报错: {e}")
        return jsonify({'error': str(e)}), 500


@app.route('/test_db')
def test_db():
    try:
        # conn = pymysql.connect(
        #     host='localhost',
        #     user='root',
        #     password='',
        #     database='ai_examine',
        #     charset='utf8mb4',
        #     unix_socket='/tmp/mysql.sock',
        #     cursorclass=pymysql.cursors.DictCursor
        # )
        conn = get_connection()
        print("✅ MySQL connection:", conn)
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM ai_examine")
        result = cur.fetchone()
        return f"数据库连接正常，共有 {result['COUNT(*)']} 条记录"
    except Exception as e:
        return f"数据库连接失败: {e}"

@app.route('/debug')
def debug_route():
    print("✅ 已进入 /debug 路由")
    return "Hello"

@app.route('/export/auditfile/<int:file_id>', methods=['GET'])
def export_single_auditfile(file_id):
    try:
        print(f"⚙️ 正在导出 file_id = {file_id}")
        data = get_news_data(file_id)
        print(f"📦 获取到数据条数：{len(data)}")

        if not data:
            return {"error": f"没有找到ID为 {file_id} 的数据"}, 404

        for i, row in enumerate(data):
            print(f"[DEBUG] 第 {i+1} 条记录：")
            print(f"  标题：{row.get('title')}")
            print(f"  reply 长度：{len(row.get('reply', ''))}")
            print(f"  reply 内容预览：{row.get('reply', '')[:100]}")

        word_file = generate_news_document(data)
        return send_file(
            word_file,
            as_attachment=True,
            download_name=f"AI辅助审查报告表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    except Exception as e:
        print(f"❌ 报错：{e}")
        return {"error": str(e)}, 500
@app.route('/test/save_word')
def test_save_word():
    from docx import Document
    from datetime import datetime

    # 假设写死两条模拟数据
    test_data = [
        {"title": "测试标题 1", "reply": "这是第一段内容。\n包含多行文字。\n确保能正确显示。"},
        {"title": "测试标题 2", "reply": "这是第二段内容。\n\n注意空行是否也能处理。"}
    ]

    # 第一个文档用你的原来样式尝试写入
    doc1 = Document()
    for d in test_data:
        doc1.add_heading(d['title'], level=2)
        doc1.add_paragraph(d['reply'])  # 不加样式
    doc1.save("with_style.docx")

    # 第二个文档用最基础方法写入，不加任何样式
    doc2 = Document()
    for d in test_data:
        doc2.add_paragraph(d['title'])
        doc2.add_paragraph(d['reply'])
    doc2.save("no_style.docx")

    return "✅ Word 文档生成完毕！请查看项目目录下的 with_style.docx 和 no_style.docx 文件。"

# 获取全部
@app.route('/api/info/list', methods=['GET'])
def get_ai_examine_files():
    """查询 ai_examine 表中的 file_id 和 file_name"""
    try:
            # 获取结果
            data = get_news_data()
            return jsonify({
                "status": "success",
                "data": data,
                "count": len(data)
            })
            
    except Exception as e:
        return jsonify({
            "status": "error",
            "message": f"数据库查询失败: {str(e)}"
        }), 500

if __name__ == '__main__':
    print("🧪 正在测试数据库连接...")
    with app.app_context():  # ✅ 创建 Flask 应用上下文
        try:
            conn = get_connection()
            with conn.cursor() as cursor:
                cursor.execute("SELECT 1")
            print("✅ 数据库连接成功！")
        except Exception as e:
            print(f"❌ 数据库连接失败: {e}")


    app.run(host='0.0.0.0', port=5005, debug=True)